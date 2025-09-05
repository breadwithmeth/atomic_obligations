#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse, json, re, sys, os, datetime
from collections import defaultdict, Counter
from typing import Dict, List, Any, Tuple

# ====== I/O: чтение PDF/DOCX/TXT ======
def read_text(path: str) -> str:
    ext = os.path.splitext(path.lower())[1]
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    elif ext == ".pdf":
        import pdfplumber
        text = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text.append(page.extract_text() or "")
        return "\n".join(text)
    else:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

def norm_text(s: str) -> str:
    s = s.replace("\u00a0", " ")  # NBSP
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()

# ====== Хелперы: поиск дат, кодов МКБ, номеров бокса/палаты ======
DATE_RE = re.compile(r"\b(\d{2})\.(\d{2})\.(\d{4})\b")
ICD_RE  = re.compile(r"\b([A-TV-ZА-ЯЁ]{1}\d{2}(?:\.\d)?)\b", re.IGNORECASE)
BOX_RE  = re.compile(r"(?:Бокс|Палата)\s*№\s*([0-9]+)")

def extract_dates(text: str) -> List[datetime.date]:
    res = []
    for d, m, y in DATE_RE.findall(text):
        try:
            res.append(datetime.date(int(y), int(m), int(d)))
        except ValueError:
            pass
    return res

def extract_icd(text: str) -> List[str]:
    # грубый сбор кодов
    return list({m.group(1).upper().replace("Ё","Е") for m in ICD_RE.finditer(text)})

def extract_boxes(text: str) -> List[str]:
    return list({m.group(1) for m in BOX_RE.finditer(text)})

# ====== Грубая нарезка на разделы (регексы можно расширять) ======
SECTION_PATTERNS = [
    ("Титул/шапка", r"(?:^|\n)(?:Медицинская документация|СТАЦИОНАРЛЫҚ|ФОРМА\s*№\s*001|ИИН|Фамилия)"),
    ("Триаж", r"(?:^|\n)Триаж"),
    ("Осмотр приёмного покоя", r"(?:^|\n)Осмотр врача приемного покоя"),
    ("Первичный осмотр", r"(?:^|\n)Первичный\s+осмотр"),
    ("Сестринский осмотр", r"(?:^|\n)Первичный\s+сестринский\s+осмотр"),
    ("Лист назначений", r"(?:^|\n)Лист(?:\s+врачебных)?\s+назначений"),
    ("Назначения на исследования", r"(?:^|\n)Лист\s+назначений\s+на\s+исследование|ЗЕРТТЕУГЕ ТАҒАЙЫНДАУЛАР ПАРАҒЫ"),
    ("Температурный лист / показатели", r"(?:^|\n)ТЕМПЕРАТУРНЫЙ ЛИСТ|ПОКАЗАТЕЛИ ЗДОРОВЬЯ"),
    ("Результаты исследований", r"(?:^|\n)Результаты исследований"),
    ("Обоснование диагноза/эпикриз", r"(?:^|\n)(Обоснование диагноза|Выписн|Эпикриз)"),
]

def split_sections(text: str) -> Dict[str, str]:
    # создаём индекс начала каждого паттерна
    marks = []
    for name, pat in SECTION_PATTERNS:
        m = re.search(pat, text, flags=re.IGNORECASE)
        if m:
            marks.append((m.start(), name))
    marks.sort()
    if not marks:
        return {"Документ целиком": text}

    # разрезаем по найденным якорям
    sections = {}
    for i, (start, name) in enumerate(marks):
        end = marks[i+1][0] if i+1 < len(marks) else len(text)
        chunk = text[start:end].strip()
        if chunk:
            sections[name] = chunk
    return sections

# ====== Правила "детерминированных" проверок (без LLM) ======
def deterministic_checks(full_text: str, sections: Dict[str,str]) -> List[Dict[str,Any]]:
    issues = []

    # 1) Скачки года в датах (например, 2024 vs 2025)
    years = re.findall(r"\b20\d{2}\b", full_text)
    count = Counter(years)
    if len(count) >= 2:
        common = count.most_common()
        dominant_year = common[0][0]
        others = [y for y,_ in common[1:]]
        issues.append({
            "severity":"critical",
            "section":"Документ",
            "title":"Несоответствие годов в датах",
            "evidence":f"Встречаются разные годы: {dict(count)}",
            "fix":"Привести все даты к фактическому году госпитализации (исправить опечатки в анамнезе/осмотрах)."
        })

    # 2) Несоответствие диагноза МКБ для тонзиллита/паратонзиллита (пример-эвристика)
    icds = extract_icd(full_text)
    if any("лакунарн" in full_text.lower() or "тонзиллит" in full_text.lower() for _ in [0]):
        wrong_b = any(code.startswith("B") for code in icds)
        if wrong_b:
            issues.append({
                "severity":"critical",
                "section":"Диагноз",
                "title":"МКБ-10 не соответствует тонзиллиту/паратонзиллиту",
                "evidence":f"Обнаружены коды: {icds}",
                "fix":"Для острого тонзиллита — J03.x; для паратонзиллярного абсцесса/целлюлита — J36. Уточнить и проставить корректные коды."
            })

    # 3) Разные номера бокса/палаты без записи о переводе
    boxes = extract_boxes(full_text)
    if len(boxes) > 1:
        issues.append({
            "severity":"major",
            "section":"Размещение",
            "title":"Разные номера бокса/палаты",
            "evidence":f"Встречаются номера: {boxes}",
            "fix":"Унифицировать место; при переводах — отдельная запись с датой/временем и основанием."
        })

    # 4) Изоляция отмечена/нет в триаже (эвристика по ключевым словам)
    triage = sections.get("Триаж","")
    if triage:
        if "Пациент должен быть изолирован" in triage and not re.search(r"изолирован:\s*(Да|Нет)", triage, re.IGNORECASE):
            issues.append({
                "severity":"major",
                "section":"Триаж",
                "title":"Не отмечено решение по изоляции",
                "evidence":"В форме триажа есть поле, но нет выбранного «Да/Нет».",
                "fix":"Отметить «Да/Нет» и указать режим изоляции (контактный/капельный и т.п.)."
            })

    # 5) План ↔ лист назначений (грубая сверка: цефтриаксон vs амокси/клав)
    plan = sections.get("Осмотр приёмного покоя","") + "\n" + sections.get("Первичный осмотр","")
    orders = sections.get("Лист назначений","")
    if plan and orders:
        need_cef = bool(re.search(r"\bцеф ?(3|триаксон)|ceftriax", plan, re.IGNORECASE))
        got_cef  = bool(re.search(r"\bцеф ?(3|триаксон)|ceftriax", orders, re.IGNORECASE))
        if need_cef and not got_cef:
            issues.append({
                "severity":"critical",
                "section":"Назначения",
                "title":"В плане заявлена парентеральная АБ, в листе назначений её нет",
                "evidence":"В плане — цефтриаксон; в листе назначений не найден.",
                "fix":"Добавить фактические инъекции с дозой/кратностью/временем и отметками исполнения или задокументировать смену схемы на per os (step-down)."
            })

    # 6) Результаты лаборатории упомянуты как заказы, но без протоколов
    lab_orders = sections.get("Назначения на исследования","")
    labs_results = sections.get("Результаты исследований","")
    if lab_orders and (not labs_results or len(labs_results) < 50):
        issues.append({
            "severity":"major",
            "section":"Исследования",
            "title":"Есть заказы на исследования, но нет (или минимум) результатов",
            "evidence":"В листе назначений на исследования много пунктов; раздел с результатами пуст/скуден.",
            "fix":"Подшить результаты (ОАК, биохимия, СРБ, мазки/ПЦР и пр.) и отразить интерпретацию в дневниках."
        })

    return issues

# ====== Вызов Ollama (локальный HTTP) ======
import urllib.request

def ollama_chat(model: str, system: str, user: str, json_schema_hint: str = "") -> Dict[str,Any]:
    url = "http://localhost:11434/api/chat"
    body = {
        "model": model,
        "messages": [
            {"role":"system","content": system},
            {"role":"user","content": user}
        ],
        "stream": False,
        "format": "json",
        "options": {"temperature": 0, "num_ctx": 8192}
    }
    req = urllib.request.Request(url, data=json.dumps(body).encode("utf-8"),
                                 headers={"Content-Type":"application/json"})
    with urllib.request.urlopen(req, timeout=300) as resp:
        data = json.loads(resp.read())
    # Ollama возвращает {"message":{"content":"{...json...}"}}
    try:
        content = data["message"]["content"]
        return json.loads(content)
    except Exception:
        return {"error":"bad_json_from_model","raw":data}

SECTION_PROMPT = """Ты — медицинский аудитор стационара (инфекционный профиль) в РК.
Проверь фрагмент истории болезни строго по чек-листу: прием, эпиданамнез, триаж/изоляция, дневники, назначения (дозы/кратность/подписи и отметки выполнения), исследования (что назначено, что выполнено, интерпретация), инфекционный контроль, переводы/боксы/палаты, выписка/эпикриз/рекомендации/лист нетрудоспособности.
Ищи противоречия (даты/годы, разные боксы без записи перевода, диагноз vs МКБ vs назначения, «в плане есть — в листе нет», «в результатах пусто» и т.д.).

Верни СТРОГО JSON по схеме:
{
  "section": "<название раздела>",
  "issues": [
    {
      "severity": "critical|major|minor",
      "title": "...",
      "evidence": "короткая цитата/факт из текста",
      "fix": "конкретное исправление одной фразой"
    }
  ]
}
Без пояснений и рассуждений.
"""

def audit_sections_with_llm(model: str, sections: Dict[str,str]) -> List[Dict[str,Any]]:
    results = []
    for name, chunk in sections.items():
        short = chunk if len(chunk) < 18000 else chunk[:18000]  # грубое отсечение
        user = f"Раздел: {name}\n\nТекст:\n<<<\n{short}\n>>>\n"
        res = ollama_chat(model, SECTION_PROMPT, user)
        if isinstance(res, dict) and "issues" in res:
            res["section"] = res.get("section", name)
            results.append(res)
        else:
            results.append({
                "section": name,
                "issues": [{"severity":"minor","title":"Модель вернула нечитаемый JSON","evidence":str(res)[:500],"fix":"Повторить прогон/уменьшить фрагмент"}]
            })
    return results

# ====== Выгрузки ======
def save_json_report(path: str, report: Dict[str,Any]):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

def save_xlsx_report(path: str, report: Dict[str,Any]):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.append(["Severity","Section","Title","Evidence","Fix"])
    for issue in report["issues_flat"]:
        ws.append([issue["severity"], issue["section"], issue["title"], issue["evidence"], issue["fix"]])
    wb.save(path)

def flatten_issues(det_issues, llm_blocks):
    arr = []
    for it in det_issues:
        arr.append({**it})
    for block in llm_blocks:
        sec = block.get("section","")
        for it in block.get("issues", []):
            arr.append({"section":sec, **it})
    return arr

def summarize(issues_flat):
    c = Counter(x["severity"] for x in issues_flat)
    return {"critical": c.get("critical",0), "major": c.get("major",0), "minor": c.get("minor",0), "total": sum(c.values())}

# ====== CLI ======
def main():
    ap = argparse.ArgumentParser(description="Авто-аудит истории болезни с помощью Ollama")
    ap.add_argument("input", help="Путь к .pdf/.docx/.txt")
    ap.add_argument("--model", default="llama3.1:8b", help="Модель Ollama (напр. llama3.1:8b)")
    ap.add_argument("--xlsx", action="store_true", help="Сохранить также .xlsx отчёт")
    args = ap.parse_args()

    raw = read_text(args.input)
    text = norm_text(raw)
    sections = split_sections(text)

    det = deterministic_checks(text, sections)
    llm = audit_sections_with_llm(args.model, sections)
    flat = flatten_issues(det, llm)
    summary = summarize(flat)

    out_json = os.path.splitext(args.input)[0] + ".audit.json"
    report = {
        "file": os.path.basename(args.input),
        "model": args.model,
        "summary": summary,
        "sections": list(sections.keys()),
        "deterministic_issues": det,
        "llm_by_section": llm,
        "issues_flat": flat
    }
    save_json_report(out_json, report)

    if args.xlsx:
        out_xlsx = os.path.splitext(args.input)[0] + ".audit.xlsx"
        save_xlsx_report(out_xlsx, report)

    # Краткий вывод
    print(f"\n=== {args.input} — итог ===")
    print(f"Критично: {summary['critical']} | Существенно: {summary['major']} | Минорно: {summary['minor']} | Всего: {summary['total']}")
    print(f"JSON: {out_json}")
    if args.xlsx:
        print(f"XLSX: {out_xlsx}")

if __name__ == "__main__":
    main()
